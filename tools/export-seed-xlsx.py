"""
Dumps the seeded tables to data/seed-data.xlsx, one sheet each, so the team can read
what is in the database without going through INSERT statements.

    python tools/export-seed-xlsx.py

Reads data/tripsense.db after the app has seeded it, rather than parsing the .sql files.
Less code, and it shows what actually landed instead of what we think the sql says. Run
it again whenever the seed changes.

Needs openpyxl: pip install -r tools/requirements.txt
"""

import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "tripsense.db"
OUT = ROOT / "data" / "seed-data.xlsx"

# sheet name -> query. joins resolve foreign keys to names, nobody wants to read ids
SHEETS = {
    "Users": "SELECT id, username, email, role FROM user ORDER BY id",
    "Countries": "SELECT id, name FROM country ORDER BY id",
    "Itineraries": """
        SELECT i.id, u.username AS user, c.name AS country, i.name,
               date(i.date_start) AS starts, date(i.date_end) AS ends
          FROM itinerary i JOIN user u ON u.id = i.user_id
          LEFT JOIN country c ON c.id = i.country_id ORDER BY i.id""",
    "Activities": """
        SELECT a.id, i.name AS itinerary, a.name, a.category, a.address,
               datetime(a.date_start) AS starts
          FROM activity a JOIN itinerary i ON i.id = a.itinerary_id ORDER BY a.id""",
    "News": "SELECT id, title, source, date(published_at) AS published FROM news ORDER BY id",
    "Impacts": """
        SELECT p.id, u.username AS user, i.name AS itinerary,
               COALESCE(a.name, '(whole itinerary)') AS activity,
               CASE p.impacted WHEN 1 THEN 'yes' ELSE 'no' END AS impacted,
               n.title AS news, p.reason
          FROM impact p JOIN user u ON u.id = p.user_id
          JOIN news n ON n.id = p.news_id
          LEFT JOIN itinerary i ON i.id = p.itinerary_id
          LEFT JOIN activity a ON a.id = p.activity_id ORDER BY p.id""",
    "ActivityOptions": """
        SELECT id, country, area, category, name, address, latitude, longitude
          FROM activity_option ORDER BY country, area, category, id""",
    "Recommendations": """
        SELECT r.id, u.username AS user, i.name AS itinerary,
               COALESCE(a.name, '(whole itinerary)') AS activity,
               r.severity, r.action, r.advice,
               o.name AS suggested, o.area AS suggested_area
          FROM recommendation r JOIN impact p ON p.id = r.impact_id
          JOIN user u ON u.id = p.user_id
          LEFT JOIN itinerary i ON i.id = p.itinerary_id
          LEFT JOIN activity a ON a.id = p.activity_id
          LEFT JOIN activity_option o ON o.id = r.suggested_option_id
          ORDER BY r.id""",
}

HEADER_FILL = PatternFill("solid", fgColor="2F5496")
MAX_WIDTH = 60


def add_sheet(book, title, rows, columns):
    sheet = book.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append(list(row))

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    for index, column in enumerate(columns, start=1):
        longest = max([len(str(column))] + [len(str(r[index - 1] or "")) for r in rows] or [0])
        sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, MAX_WIDTH)

    sheet.freeze_panes = "A2"               # headers stay put when scrolling
    sheet.auto_filter.ref = sheet.dimensions


def main():
    if not DB.exists():
        sys.exit(f"no database at {DB}, start the app once so it seeds")

    db = sqlite3.connect(DB)
    book = Workbook()
    book.remove(book.active)                # openpyxl makes an empty default sheet

    for title, query in SHEETS.items():
        try:
            cursor = db.execute(query)
        except sqlite3.OperationalError as e:
            # a table that does not exist yet, eg recommendation before part 5 lands
            print(f"  skipped {title}: {e}")
            continue
        rows = cursor.fetchall()
        add_sheet(book, title, rows, [d[0] for d in cursor.description])
        print(f"  {title:<18} {len(rows)} rows")

    OUT.parent.mkdir(exist_ok=True)
    book.save(OUT)
    print(f"\nwrote {OUT}")


if __name__ == "__main__":
    main()
